"""Core-function checks; no desktop window or CustomTkinter execution."""
import ast
from datetime import datetime
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


SOURCE = Path(__file__).resolve().parents[1] / 'app.py'


def load_core():
    source = SOURCE.read_text(encoding='utf-8')
    compile(source, str(SOURCE), 'exec')
    tree = ast.parse(source)
    data = {'properties_database', 'admin_accounts', 'user_accounts'}
    nodes = [n for n in tree.body if isinstance(n, ast.FunctionDef) or
             (isinstance(n, ast.Assign) and any(isinstance(t, ast.Name) and
              t.id in data for t in n.targets))]
    namespace = dict(datetime=datetime, openpyxl=openpyxl, Font=Font,
                     PatternFill=PatternFill, Alignment=Alignment,
                     Border=Border, Side=Side)
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(SOURCE), 'exec'), namespace)
    return namespace


class CoreChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_root = SOURCE.parent / '.test-tmp'
        cls.temp_root.mkdir(exist_ok=True)

    def setUp(self):
        self.core = load_core()
        self.data = self.core['properties_database']

    def call(self, name, *args):
        return self.core[name](*args)

    def ids(self, values):
        return [p['id'] for p in values]

    def test_type_case_and_unknown(self):
        self.assertEqual(self.ids(self.call('search_by_type', 'sToRe')), [1, 5])
        self.assertEqual(self.call('search_by_type', 'Unknown'), [])

    def test_budget_inclusive_boundary(self):
        self.assertEqual(self.ids(self.call('search_by_budget', 5000)), [1, 5])
        self.assertEqual(self.call('search_by_budget', 0), [])

    def test_size_inclusive_boundary(self):
        self.assertEqual(self.ids(self.call('search_by_size', 2000)), [2, 7])

    def test_parking_inclusive_boundary(self):
        self.assertEqual(self.ids(self.call('search_by_parking', 50)), [2, 7])

    def test_floor_equality(self):
        self.assertEqual(self.ids(self.call('search_by_floors', 3)), [6, 8])

    def test_availability_exact_matching_limitation(self):
        self.assertEqual(self.ids(self.call('search_by_availability', 'lease')), [1, 4, 8])

    def test_rating_threshold(self):
        self.assertEqual(self.ids(self.call('search_by_rating', 4.8)), [2, 6])

    def test_recommendations_are_rating_sort(self):
        self.assertEqual(self.ids(self.call('get_recommended_properties')), [6, 2, 4, 7, 1])

    def test_recommendations_budget_after_top_five_limitation(self):
        self.assertEqual(self.ids(self.call('get_recommended_properties', 5000)), [1])

    def test_add_property_uses_new_id_and_updates_search(self):
        added = self.call('add_new_property', 'Test office', 'Office', 2, 250, 3, 'Lease', 4000)
        self.assertEqual(added['id'], 9)
        self.assertIn(added, self.call('search_by_budget', 4000))
        self.assertEqual(added['rating'], 0)

    def test_fresh_load_resets_runtime_data(self):
        self.call('add_new_property', 'Temporary', 'Office', 1, 100, 1, 'Lease', 100)
        self.assertEqual(len(load_core()['properties_database']), 8)

    def test_export_round_trip(self):
        with tempfile.TemporaryDirectory(dir=self.temp_root) as folder:
            path = Path(folder) / 'report.xlsx'
            self.core['filedialog'] = SimpleNamespace(asksaveasfilename=lambda **kwargs: str(path))
            subset = self.call('search_by_budget', 5000)
            self.assertEqual(self.call('generate_excel_report', subset, 'Budget <= 5000'), str(path))
            book = openpyxl.load_workbook(path)
            sheet = book.active
            self.assertEqual(sheet['A4'].value, 'Total Results: 2')
            self.assertEqual(sheet['A3'].value, 'Search Criteria: Budget <= 5000')
            self.assertEqual([sheet.cell(r, 1).value for r in [7, 8]], [1, 5])
            self.assertEqual(sheet['H7'].value, 5000)
            self.assertEqual(sheet['H7'].number_format, '£#,##0.00')
            self.assertTrue(sheet['A6'].font.bold)
            book.close()

    def test_export_empty_results(self):
        with tempfile.TemporaryDirectory(dir=self.temp_root) as folder:
            path = Path(folder) / 'empty.xlsx'
            self.core['filedialog'] = SimpleNamespace(asksaveasfilename=lambda **kwargs: str(path))
            self.call('generate_excel_report', [], 'No match')
            book = openpyxl.load_workbook(path)
            self.assertEqual(book.active['A4'].value, 'Total Results: 0')
            self.assertIsNone(book.active['A7'].value)
            book.close()

    def test_cancelled_export(self):
        self.core['filedialog'] = SimpleNamespace(asksaveasfilename=lambda **kwargs: '')
        self.assertIsNone(self.call('generate_excel_report', self.data, 'All'))


if __name__ == '__main__':
    unittest.main()
